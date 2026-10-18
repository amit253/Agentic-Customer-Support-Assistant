"""Build the submission deck from verified project evidence.

Every number on a slide comes from a committed report or evidence file listed
in ``docs/presentation-speaker-notes.md``. The deck is generated, not
hand-edited, so it can be rebuilt after any change:

    python scripts/build_presentation.py

Design: 16:9, Cambria headings with Calibri body, deep slate-teal as the
dominant color with one amber accent, numbered amber badges as the repeated
motif, and native PowerPoint charts. Slides carry no source lines or speaker
notes so the file can be submitted as is; presenter notes and sources are
written to ``docs/presentation-speaker-notes.md``.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "agentic-customer-support.pptx"
NOTES_OUTPUT = ROOT / "docs" / "presentation-speaker-notes.md"

HEAD = "Cambria"
BODY = "Calibri"
INK = RGBColor(0x12, 0x33, 0x3B)
INK_SOFT = RGBColor(0x1D, 0x46, 0x51)
TEAL = RGBColor(0x1F, 0x7A, 0x80)
TEAL_MID = RGBColor(0x7F, 0xB3, 0xB5)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
MIST = RGBColor(0xEC, 0xF3, 0xF3)
PANEL = RGBColor(0xF6, 0xFA, 0xFA)
LINE = RGBColor(0xB8, 0xCA, 0xCC)
GRID = RGBColor(0xDD, 0xE7, 0xE8)
MUTED = RGBColor(0x55, 0x69, 0x6D)
ICE = RGBColor(0xCF, 0xE3, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
RED = RGBColor(0xB5, 0x46, 0x33)
SLATE = RGBColor(0xA9, 0xBC, 0xBE)

FOOTER = "Northstar Gear · Agentic Customer Support Assistant"
LEFT = PP_ALIGN.LEFT
CENTER = PP_ALIGN.CENTER
MIDDLE = MSO_ANCHOR.MIDDLE
TOP = MSO_ANCHOR.TOP

# Presenter notes and sources, collected per slide and written to Markdown.
SPEAKER_NOTES: list[tuple[str, str, str]] = []


# --------------------------------------------------------------------------
# Drawing helpers
# --------------------------------------------------------------------------

def _runs(paragraph, content, *, size, color, bold, font, italic=False) -> None:
    parts = content if isinstance(content, list) else [(content, {})]
    for part in parts:
        text_value, overrides = (part, {}) if isinstance(part, str) else part
        run = paragraph.add_run()
        run.text = text_value
        run.font.name = overrides.get("font", font)
        run.font.size = Pt(overrides.get("size", size))
        run.font.bold = overrides.get("bold", bold)
        run.font.italic = overrides.get("italic", italic)
        run.font.color.rgb = overrides.get("color", color)


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paragraphs,
    *,
    size: float = 14,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = BODY,
    align=LEFT,
    anchor=TOP,
    space_after: float = 0,
    italic: bool = False,
    bullets: bool = False,
    name: str | None = None,
):
    """Text box with zero padding; ``paragraphs`` is a str or a list of them."""

    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    items = paragraphs if isinstance(paragraphs, list) else [paragraphs]
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        if bullets:
            _bullet(paragraph)
        _runs(paragraph, item, size=size, color=color, bold=bold, font=font, italic=italic)
    return shape


def _bullet(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(Inches(0.22)))
    properties.set("indent", str(-Inches(0.22)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for existing in properties.findall(qn(tag)):
            properties.remove(existing)
    font = etree.SubElement(properties, qn("a:buFont"))
    font.set("typeface", "Arial")
    char = etree.SubElement(properties, qn("a:buChar"))
    char.set("char", "•")


def box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor | None = MIST,
    line: RGBColor | None = None,
    radius: float = 0.08,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    dashed: bool = False,
):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
        if dashed:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def node(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    heading: str,
    sub: str | None = None,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor | None = LINE,
    title_color: RGBColor = INK,
    sub_color: RGBColor = MUTED,
    title_size: float = 13,
    sub_size: float = 10.5,
    align=CENTER,
):
    """Diagram node: rounded box with a bold title and an optional caption."""

    shp = box(slide, x, y, w, h, fill=fill, line=line, radius=0.12)
    frame = shp.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = MIDDLE
    first = frame.paragraphs[0]
    first.alignment = align
    _runs(first, heading, size=title_size, color=title_color, bold=True, font=BODY)
    if sub:
        second = frame.add_paragraph()
        second.alignment = align
        second.space_before = Pt(2)
        _runs(second, sub, size=sub_size, color=sub_color, bold=False, font=BODY)
    return shp


def badge(slide, x: float, y: float, label: str, *, d: float = 0.42, fill=AMBER, color=INK):
    shp = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    frame = shp.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = CENTER
    _runs(paragraph, label, size=13 if d >= 0.4 else 11, color=color, bold=True, font=BODY)
    return shp


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: RGBColor = TEAL,
    width: float = 1.75,
    dashed: bool = False,
    head: bool = True,
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        line = connector.line._get_or_add_ln()
        tail = etree.SubElement(line, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return connector


def stat(slide, x: float, y: float, w: float, value: str, label: str, *, detail: str | None = None,
         value_color: RGBColor = TEAL, label_color: RGBColor = INK, size: float = 40) -> None:
    text(slide, x, y, w, 0.75, value, size=size, color=value_color, bold=True, font=HEAD)
    text(slide, x, y + 0.8, w, 0.6, label, size=14, color=label_color, bold=True)
    if detail:
        text(slide, x, y + 1.42, w, 0.5, detail, size=12, color=MUTED)


def chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    categories: Sequence[str],
    series: Sequence[tuple[str, Sequence[float], RGBColor]],
    *,
    number_format: str = "0%",
    axis_format: str | None = None,
    maximum: float | None = None,
    major_unit: float | None = None,
    horizontal: bool = False,
    legend: bool = True,
    point_colors: Sequence[RGBColor] | None = None,
):
    data = CategoryChartData()
    data.categories = list(categories)
    for name, values, _ in series:
        data.add_series(name, list(values))
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    figure = frame.chart
    figure.font.name = BODY
    figure.font.size = Pt(11)
    figure.font.color.rgb = MUTED
    figure.has_title = False
    figure.has_legend = legend and len(series) > 1
    if figure.has_legend:
        figure.legend.position = XL_LEGEND_POSITION.TOP
        figure.legend.include_in_layout = False
        figure.legend.font.size = Pt(11)
        figure.legend.font.color.rgb = INK
    plot = figure.plots[0]
    plot.gap_width = 55
    plot.overlap = -8
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(11)
    labels.font.color.rgb = INK
    for index, (_, _, color) in enumerate(series):
        fill = plot.series[index].format.fill
        fill.solid()
        fill.fore_color.rgb = color
    if point_colors:
        for index, color in enumerate(point_colors):
            point = plot.series[0].points[index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = color
    values = figure.value_axis
    values.has_major_gridlines = True
    values.major_gridlines.format.line.color.rgb = GRID
    values.major_gridlines.format.line.width = Pt(0.75)
    values.format.line.fill.background()
    values.tick_labels.font.size = Pt(10)
    values.tick_labels.font.color.rgb = MUTED
    values.tick_labels.number_format = axis_format or number_format
    values.tick_labels.number_format_is_linked = False
    values.minimum_scale = 0
    if maximum is not None:
        values.maximum_scale = maximum
    if major_unit is not None:
        values.major_unit = major_unit
    categories_axis = figure.category_axis
    categories_axis.format.line.color.rgb = LINE
    categories_axis.has_major_gridlines = False
    categories_axis.tick_labels.font.size = Pt(11)
    categories_axis.tick_labels.font.color.rgb = INK
    if horizontal:
        categories_axis.reverse_order = True
    return figure


def table(slide, x: float, y: float, w: float, header: Sequence[str], rows: Sequence[Sequence],
          widths: Sequence[float], *, size: float = 12, row_h: float = 0.42) -> None:
    shape = slide.shapes.add_table(
        len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1))
    )
    grid = shape.table
    grid.first_row = True
    grid.horz_banding = False
    for column, width in enumerate(widths):
        grid.columns[column].width = Inches(width)
    for row_index in range(len(rows) + 1):
        grid.rows[row_index].height = Inches(row_h)
        for column in range(len(header)):
            cell = grid.cell(row_index, column)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = INK
                value, color, bold = header[column], WHITE, True
            else:
                cell.fill.fore_color.rgb = WHITE if row_index % 2 else MIST
                raw = rows[row_index - 1][column]
                value, color, bold = (raw if isinstance(raw, tuple) else (raw, INK, False))
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = ""
            _runs(paragraph, value, size=size, color=color, bold=bold, font=BODY)


def title(slide, heading: str, kicker: str | None = None, *, dark: bool = False) -> None:
    text(slide, 0.6, 0.42, 12.1, 0.72, heading, size=34, bold=True, font=HEAD,
         color=WHITE if dark else INK, name="Title")
    if kicker:
        text(slide, 0.6, 1.16, 12.1, 0.42, kicker, size=16, color=ICE if dark else MUTED)


def new_slide(prs, *, dark: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = INK if dark else WHITE
    return slide


def finish(prs, slide, heading: str, notes: str, sources: str, *, dark: bool = False) -> None:
    """Add the brand footer and slide number; record notes and sources off-slide."""

    number = len(prs.slides._sldIdLst)
    color = ICE if dark else MUTED
    text(slide, 0.6, 7.02, 8.0, 0.25, FOOTER, size=10, color=color)
    text(slide, 11.9, 7.02, 0.85, 0.25, str(number), size=10, color=color, align=PP_ALIGN.RIGHT)
    SPEAKER_NOTES.append((heading, notes, sources))


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------

def slide_cover(prs) -> None:
    slide = new_slide(prs, dark=True)
    star = box(slide, 0.62, 1.06, 0.34, 0.34, fill=AMBER, shape=MSO_SHAPE.STAR_4_POINT)
    star.name = "North star"
    text(slide, 1.08, 1.08, 10, 0.32, "AI / LLM ENGINEERING ASSIGNMENT  ·  FINAL SUBMISSION",
         size=13, color=AMBER, bold=True)
    heading = "Agentic Customer Support Assistant"
    text(slide, 0.6, 1.7, 12.0, 0.95, heading, size=46, bold=True, font=HEAD, color=WHITE, name="Title")
    text(slide, 0.6, 2.75, 11.4, 0.9,
         "Text and voice support for Northstar Gear, a fictional travel-backpack brand, "
         "with Groq for LLM and speech, grounded answers, read-only tools, and guardrails in code",
         size=19, color=ICE)
    chips = [
        ("95.5%", "routing accuracy, live, on 110 held-out messages"),
        ("13 / 13", "unanswerable questions declined; 93.8% of answerable ones answered"),
        ("2.1–2.9 s", "voice: end of speech to start of the spoken reply"),
    ]
    for index, (value, label) in enumerate(chips):
        x = 0.6 + index * 4.1
        box(slide, x, 4.55, 3.8, 1.55, fill=INK_SOFT, radius=0.1)
        text(slide, x + 0.3, 4.72, 3.3, 0.65, value, size=34, bold=True, font=HEAD, color=AMBER)
        text(slide, x + 0.3, 5.42, 3.3, 0.6, label, size=13, color=WHITE)
    finish(prs, slide, heading,
           "One sentence: a support assistant that works over text and a phone-style voice line, where "
           "every safety property is enforced in application code and every number comes from a "
           "committed report.",
           "reports/phase2-hybrid-live-low-effort-evaluation.md; reports/final-rag-live-evaluation.md; "
           "recordings/final-*-run.json", dark=True)


def slide_problem(prs) -> None:
    slide = new_slide(prs)
    heading = "The problem: support over text and voice"
    title(slide, heading,
          "Build a customer-support assistant on Groq that routes, grounds, calls tools, refuses "
          "unsafe requests, and behaves like a phone line")
    parts = [
        ("A", "Agentic support system", [
            "Route messages to 3+ intents, each with its own agent",
            "RAG agent grounded in a 10–20 document knowledge base",
            "Tool-calling agent for order status and return eligibility",
            "Input guard for prompt injection; output guard for forbidden claims",
            "Cache classifier calls and report the hit rate",
        ]),
        ("B", "Voice turn-taking pipeline", [
            "Audio in, Groq Whisper STT, TTS out, framed as a phone line",
            "Barge-in: silence, noise burst, real interruption, own adversarial case",
            "Provider failure mid-call: retry or fall back, never hang",
            "Measure end of speech to the first audio of the reply",
        ]),
        ("C", "Integration and operations", [
            "Share state when a call moves to text chat or a human",
            "Production evaluation: guardrail regression, latency SLOs, routing A/B",
            "Groq degradation: detection, fallback, and failover trade-offs",
        ]),
    ]
    for index, (letter, name, items) in enumerate(parts):
        x = 0.6 + index * 4.1
        box(slide, x, 1.85, 3.85, 3.65, fill=PANEL, line=GRID)
        badge(slide, x + 0.25, 2.05, letter)
        text(slide, x + 0.8, 2.1, 2.9, 0.35, name, size=16, bold=True)
        text(slide, x + 0.3, 2.7, 3.3, 2.75, items, size=13, bullets=True, space_after=5)
    text(slide, 0.6, 5.72, 3.0, 0.3, "Constraints that shape the design", size=14, bold=True, color=TEAL)
    constraints = ["Own domain and knowledge base", "Real recorded audio", "Grounding, not guessing",
                   "Guardrails hold under pressure", "No dropped or duplicated turns", "Keys never committed"]
    for index, label in enumerate(constraints):
        x = 0.6 + index * 2.05
        node(slide, x, 6.1, 1.9, 0.6, label, fill=MIST, line=None, title_size=11.5)
    finish(prs, slide, heading,
           "The assignment in one slide. Part A is the text agent system, Part B the phone-style voice "
           "line, Part C written integration and operations answers. The constraint strip is what "
           "reviewers test hardest: grounding, guardrails under pressure, and no dropped or duplicated "
           "voice turns.",
           "Assignment problem statement (kept locally under Problem_Statement/, not in the repository)")


def slide_summary(prs) -> None:
    slide = new_slide(prs)
    heading = "Solution at a glance"
    title(slide, heading, "Measured on held-out sets with live Groq calls; every figure is reproducible")
    stats = [
        ("95.5%", "Routing accuracy", "110 held-out messages, live"),
        ("93.8%", "Answerable questions answered", "13 / 13 unanswerable declined"),
        ("0.922", "Evidence recall", "expected documents reaching the answerer"),
        ("2.1–2.9 s", "Voice response start", "end of speech to spoken reply"),
    ]
    for index, (value, label, detail) in enumerate(stats):
        x = 0.6 + index * 3.1
        box(slide, x, 1.85, 2.85, 2.1, fill=MIST)
        stat(slide, x + 0.25, 2.0, 2.4, value, label, detail=detail, size=36)
    steps = [
        ("1", "Hybrid router", "Rules, then labeled examples, then the LLM only when needed: 0.245 LLM calls per message"),
        ("2", "Conversation agent", "Greetings, off-topic and unclear input get a fixed, safe reply, never a dead end"),
        ("3", "Hybrid retrieval", "e5 + BM25 fusion; the gate favors recall and the answerer declines without evidence"),
        ("4", "Voice built per call", "Models load when the call connects, so each turn only pays for STT, routing and TTS"),
    ]
    for index, (number, name, body) in enumerate(steps):
        x = 0.6 + (index % 2) * 6.2
        y = 4.3 + (index // 2) * 1.15
        badge(slide, x, y + 0.02, number)
        text(slide, x + 0.6, y, 5.4, 0.32, name, size=16, bold=True)
        text(slide, x + 0.6, y + 0.34, 5.4, 0.65, body, size=14, color=MUTED)
    finish(prs, slide, heading,
           "Four headline numbers and the four design moves behind them. All live figures used Groq "
           "openai/gpt-oss-20b on held-out sets committed before tuning.",
           "reports/phase2-hybrid-live-low-effort-evaluation.md; reports/final-rag-live-evaluation.md; "
           "reports/final-retrieval-offline-evaluation.md; recordings/final-*-run.json")


def slide_requirements(prs) -> None:
    slide = new_slide(prs)
    heading = "Every required capability, with its status"
    title(slide, heading, "✓ implemented and evidenced   ◐ partial, and stated as partial")
    columns = [
        ("A", "Part A · Agentic support", [
            ("✓", "Router: 6 intents, distinct handlers"),
            ("✓", "RAG over 12 KB docs with embeddings"),
            ("✓", "Groq native function calling, read-only tools"),
            ("✓", "Input and output guards in code, fail-closed"),
            ("✓", "Classifier cache with measured hit rate"),
            ("✓", "Refund vs eligibility reconciliation"),
        ]),
        ("B", "Part B · Voice line", [
            ("✓", "Barge-in: silence, noise, speech, adversarial"),
            ("✓", "STT failure: retry, then safe spoken handoff"),
            ("✓", "Latency instrumented per turn"),
            ("✓", "Real recorded audio, not synthetic input"),
            ("◐", "First-acoustic-sample timing not measured"),
            ("◐", "Live microphone barge-in not demonstrated"),
        ]),
        ("C", "Part C and process", [
            ("✓", "Voice and text share session state + handoff"),
            ("◐", "State is process-local, not durable"),
            ("◐", "LLM provider failover designed, not built"),
            ("✓", "Incremental commits and decisions log"),
            ("✓", "Red-team findings reproduced, open ones kept"),
            ("✓", "Held-out eval sets, live vs offline labels"),
        ]),
    ]
    for index, (letter, name, items) in enumerate(columns):
        x = 0.6 + index * 4.1
        box(slide, x, 1.85, 3.85, 4.75, fill=PANEL, line=GRID)
        badge(slide, x + 0.25, 2.05, letter)
        text(slide, x + 0.8, 2.1, 2.9, 0.35, name, size=16, bold=True)
        for row, (mark, label) in enumerate(items):
            y = 2.75 + row * 0.62
            text(slide, x + 0.3, y, 0.3, 0.35, mark, size=16, bold=True,
                 color=GREEN if mark == "✓" else AMBER)
            text(slide, x + 0.65, y, 3.05, 0.55, label, size=14)
    finish(prs, slide, heading,
           "The partial items are deliberate honesty: first-acoustic-sample latency needs a loopback rig; "
           "durable memory and provider failover are designed in the docs but not built for the demo.",
           "docs/requirements-matrix.md (Final status); docs/design-notes.md; docs/part-c-answers.md")


def slide_architecture(prs) -> None:
    slide = new_slide(prs)
    heading = "Architecture: one guarded service for text and voice"
    title(slide, heading)
    node(slide, 0.6, 1.85, 2.2, 0.95, "Text chat", "Streamlit UI, mic in the box")
    node(slide, 0.6, 3.05, 2.2, 1.75, "Voice line",
         "WebRTC + Silero VAD, barge-in coordinator, Groq Whisper STT")
    node(slide, 0.6, 5.1, 2.2, 0.95, "Reply", "text, or local TTS playback")
    container = box(slide, 3.15, 1.75, 6.95, 4.65, fill=PANEL, line=LINE, radius=0.04)
    container.name = "SupportService"
    text(slide, 3.4, 1.88, 6.4, 0.3, "SupportService  ·  shared by every channel", size=12, bold=True, color=TEAL)
    row1 = [("Validate", "NFKC, size, control chars"), ("Input guard", "rules + Groq classifier"),
            ("Hybrid router", "rules → examples → LLM")]
    for index, (name, sub) in enumerate(row1):
        node(slide, 3.4 + index * 2.25, 2.3, 1.95, 0.85, name, sub)
    for index in range(2):
        arrow(slide, 5.35 + index * 2.25, 2.72, 5.65 + index * 2.25, 2.72)
    agents = [("RAG agent", "e5 + BM25, cited"), ("Order agent", "read-only tools"),
              ("Handoff", "summary + reference"), ("Conversation", "templates only")]
    arrow(slide, 8.875, 3.15, 8.875, 3.42, head=False)
    arrow(slide, 4.125, 3.42, 9.135, 3.42, head=False)
    for index, (name, sub) in enumerate(agents):
        x = 3.4 + index * 1.67
        arrow(slide, x + 0.725, 3.42, x + 0.725, 3.68)
        node(slide, x, 3.68, 1.45, 1.0, name, sub, title_size=12.5, sub_size=10)
        arrow(slide, x + 0.725, 4.68, x + 0.725, 4.95, head=False)
    arrow(slide, 4.125, 4.95, 9.135, 4.95, head=False)
    arrow(slide, 6.6, 4.95, 6.6, 5.25)
    node(slide, 4.6, 5.25, 4.0, 0.9, "Output claim guard", "blocks unverified IDs, dates, amounts, promises",
         fill=MIST)
    arrow(slide, 2.8, 2.32, 3.4, 2.62)
    arrow(slide, 2.8, 3.9, 3.4, 2.85)
    arrow(slide, 4.6, 5.7, 2.8, 5.58)
    deps = [
        ("Groq API", "gpt-oss-20b: guard, router tier 2, answers, tool choice · Whisper STT", 1.75, 1.15),
        ("Local models", "MiniLM routing examples · e5 retrieval", 3.05, 0.85),
        ("Knowledge base", "12 docs · hybrid index", 4.05, 0.72),
        ("Order records", "read-only · ownership-checked", 4.92, 0.72),
        ("Session memory", "verified slots + transcript", 5.79, 0.72),
    ]
    for name, sub, y, h in deps:
        node(slide, 10.45, y, 2.3, h, name, sub, fill=MIST, line=None, title_size=12, sub_size=10)
        arrow(slide, 10.45, y + h / 2, 10.1, y + h / 2, color=SLATE, width=1.25, dashed=True)
    finish(prs, slide, heading,
           "Read left to right. Both channels enter the same SupportService. Nothing reaches the router or "
           "any agent before validation and the input guard; every reply leaves through the output claim "
           "guard. Dashed lines are dependencies. Conversation agents have no retrieval or tools.",
           "src/agentic_support/pipeline.py (SupportService); hybrid_routing.py; handlers.py; voice.py")


def slide_lifecycle(prs) -> None:
    slide = new_slide(prs)
    heading = "Every step fails safe, and none of them fails silently"
    title(slide, heading, "Each stage has a named failure outcome and a customer-facing reply")
    steps = [
        ("1", "Validate", "Unicode, size, control characters", "Rejected: “couldn’t read that”"),
        ("2", "Input guard", "Rule signals + Groq safety classifier", "Refused, or “can’t check that safely right now”"),
        ("3", "Route", "Rules → examples → LLM", "Clarifying question, never a dead end"),
        ("4", "Handle", "RAG, order tools, handoff, conversation", "Stopped “rather than guess”"),
        ("5", "Output guard", "Facts must come from tools or citations", "Blocked: “couldn’t verify every detail”"),
        ("6", "Remember", "Every turn, refused ones included", "Only verified slots ever reach an LLM"),
    ]
    for index, (number, name, body, failure) in enumerate(steps):
        x = 0.6 + index * 2.05
        box(slide, x, 1.85, 1.88, 2.05, fill=MIST)
        badge(slide, x + 0.18, 2.0, number)
        text(slide, x + 0.18, 2.55, 1.6, 0.35, name, size=16, bold=True)
        text(slide, x + 0.18, 2.95, 1.6, 0.9, body, size=12, color=MUTED)
        if index < 5:
            arrow(slide, x + 1.88, 2.87, x + 2.05, 2.87)
        box(slide, x, 4.1, 1.88, 1.25, fill=WHITE, line=LINE)
        text(slide, x + 0.15, 4.2, 1.6, 0.25, "If it fails", size=11, bold=True, color=RED)
        text(slide, x + 0.15, 4.47, 1.6, 0.85, failure, size=12)
    box(slide, 0.6, 5.6, 12.15, 0.95, fill=INK)
    star = box(slide, 0.85, 5.88, 0.38, 0.38, fill=AMBER, shape=MSO_SHAPE.STAR_4_POINT)
    star.name = "Guarantee"
    text(slide, 1.45, 5.68, 11.1, 0.8,
         "Structural guarantee: the router and every agent run only after the input guard allows the message. "
         "Unrouted or unclear messages go to an agent with no retrieval and no tools.",
         size=15, color=WHITE, anchor=MIDDLE)
    finish(prs, slide, heading,
           "This is the end-to-end trace the assignment asks for. Running the guard and the router in "
           "parallel was considered and rejected because the router LLM would then see unchecked text.",
           "src/agentic_support/pipeline.py; tests/test_step5_call_path.py; tests/test_conversation.py")


def slide_router(prs) -> None:
    slide = new_slide(prs)
    heading = "Hybrid routing: the LLM only when needed"
    title(slide, heading, "Share of the 110 held-out messages each tier decided in the live run")
    tiers = [
        ("0", "Rules", "order ID + operation, verified follow-up, smalltalk, human request, gibberish", 54, "~0 ms"),
        ("1", "Labeled examples", "nearest examples on MiniLM; answers only with a clear margin", 29, "~6 ms"),
        ("2", "Groq LLM", "strict schema with all six intents; validated locally", 27, "~0.6 s"),
        ("?", "Clarify", "no confident route: a clarifying question, not a refusal", 0, "0 ms"),
    ]
    for index, (number, name, body, count, cost) in enumerate(tiers):
        y = 1.9 + index * 1.15
        badge(slide, 0.6, y + 0.1, number)
        text(slide, 1.2, y, 2.6, 0.35, name, size=16, bold=True)
        text(slide, 1.2, y + 0.38, 3.6, 0.65, body, size=12, color=MUTED)
        box(slide, 5.0, y + 0.12, 4.4, 0.42, fill=GRID, radius=0.2)
        if count:
            box(slide, 5.0, y + 0.12, max(0.15, 4.4 * count / 110), 0.42, fill=TEAL, radius=0.2)
        text(slide, 5.0, y + 0.62, 2.5, 0.3, f"{count} of 110", size=12, bold=True)
        text(slide, 7.5, y + 0.62, 1.9, 0.3, cost, size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    box(slide, 9.85, 1.85, 2.9, 4.55, fill=MIST)
    text(slide, 10.1, 2.0, 2.5, 0.35, "Why hybrid", size=16, bold=True)
    text(slide, 10.1, 2.45, 2.5, 3.9, [
        "Rules are instant and deterministic for unambiguous requests",
        "Labeled examples cover paraphrases in milliseconds, with no API call",
        "The LLM handles Hinglish, mixed and novel phrasing",
        "Thresholds come from leave-one-out on 137 training examples, never the test set",
    ], size=13, bullets=True, space_after=6)
    finish(prs, slide, heading,
           "Rules decide 54 of 110, labeled examples 29, the LLM 27. A test fails if any training example "
           "appears in the held-out set. The semantic tier answers only when the best intent beats the "
           "runner-up by 0.15 cosine; that margin gave precision 1.00 at 44% coverage in leave-one-out.",
           "reports/phase2-hybrid-live-evaluation.md (routing sources); docs/router-design.md")


def slide_routing_results(prs) -> None:
    slide = new_slide(prs)
    heading = "Routing results on 110 held-out messages"
    title(slide, heading, "Live run with Groq openai/gpt-oss-20b, final configuration")
    intents = ["Product question", "Order status", "Human handoff", "Smalltalk", "Out of scope", "Unclear"]
    chart(
        slide, 0.6, 1.75, 8.5, 4.95, intents,
        [("Precision", [1.0, 0.923, 0.933, 1.0, 0.933, 1.0], TEAL),
         ("Recall", [0.917, 1.0, 0.933, 1.0, 1.0, 0.923], SLATE)],
        number_format="0%", maximum=1.12, major_unit=0.2,
    )
    callouts = [("95.5%", "strict accuracy"), ("99.1%", "outcome the customer sees"),
                ("0.70 s", "routing p95 latency"), ("0", "false escalations to a human")]
    for index, (value, label) in enumerate(callouts):
        y = 1.85 + index * 1.2
        box(slide, 9.45, y, 3.3, 1.05, fill=MIST)
        text(slide, 9.7, y + 0.1, 2.9, 0.5, value, size=26, bold=True, font=HEAD, color=TEAL)
        text(slide, 9.7, y + 0.62, 2.9, 0.35, label, size=13, bold=True)
    finish(prs, slide, heading,
           "Per-intent precision and recall on the held-out set. 'Outcome' counts a no-route result as "
           "correct only when the clarifying reply is the right answer. Remaining misses: 'the blue one' "
           "gets a clarifying question; 'xkcd lorem ipsum' is called off-topic instead of unclear.",
           "reports/phase2-hybrid-live-low-effort-evaluation.md")


def slide_retrieval(prs) -> None:
    slide = new_slide(prs)
    heading = "Retrieval: fuse meaning and keywords, gate for recall"
    title(slide, heading, "Five configurations compared: tuned on a dev set, scored on 45 held-out questions")
    flow = [
        ("Query", None), ("e5 dense  +  BM25 keywords", "multilingual paraphrase · exact names"),
        ("Reciprocal-rank fusion", "no score-scale mixing"), ("Gate on best candidate", "tuned for recall"),
        ("Answerer", "declines if evidence is insufficient"), ("Output claim guard", "dates, amounts, IDs"),
    ]
    for index, (name, sub) in enumerate(flow):
        y = 1.8 + index * 0.82
        node(slide, 0.6, y, 3.6, 0.66, name, sub, fill=MIST if index in (3, 4) else WHITE,
             title_size=12.5, sub_size=10)
        if index < len(flow) - 1:
            arrow(slide, 2.4, y + 0.66, 2.4, y + 0.82)
    chart(
        slide, 4.6, 1.75, 8.2, 4.95,
        ["MiniLM dense", "MiniLM dense, tuned cut", "MiniLM + BM25",
         "MiniLM + BM25 + reranker", "e5 + BM25 (selected)"],
        [("Evidence recall", [0.516, 0.781, 0.797, 0.562, 0.922], TEAL),
         ("MRR", [0.857, 0.857, 0.907, 0.885, 0.912], SLATE)],
        number_format="0.00", maximum=1.0, horizontal=True,
    )
    finish(prs, slide, heading,
           "Why this embedding approach: e5 is multilingual (Hinglish) and fused with BM25 for exact product "
           "names. Cross-encoder rerankers were tried and dropped: their dev-tuned gate did not transfer and "
           "they cost about 0.5 s per query on CPU.",
           "reports/retrieval-comparison.md; reports/final-retrieval-offline-evaluation.md")


def slide_grounding(prs) -> None:
    slide = new_slide(prs)
    heading = "Grounded answers: two gates and a claim check"
    title(slide, heading, "Live end-to-end on 45 held-out questions: retrieval, Groq answerer, output guard")
    stats = [("93.8%", "answerable questions answered"), ("90.6%", "answers citing an expected document"),
             ("13 / 13", "unanswerable near-misses declined"), ("0", "handler errors")]
    for index, (value, label) in enumerate(stats):
        y = 1.9 + index * 1.15
        text(slide, 0.6, y, 2.4, 0.6, value, size=32, bold=True, font=HEAD, color=TEAL)
        text(slide, 3.05, y + 0.12, 3.0, 0.6, label, size=15, bold=True)
    box(slide, 6.5, 1.85, 6.25, 4.75, fill=MIST)
    badge(slide, 6.75, 2.05, "$")
    text(slide, 7.35, 2.08, 5.2, 0.4, "Refunds: tool facts vs claim guard", size=17, bold=True)
    text(slide, 6.8, 2.7, 5.7, 2.2, [
        "Return eligibility is policy evidence, never refund approval",
        "Only an amount and date confirmed by the order record can appear",
        "“Refund my amount” gets: I can’t issue refunds in chat, the verified facts, and a route to a human",
    ], size=14, bullets=True, space_after=8)
    box(slide, 6.8, 4.95, 5.7, 1.3, fill=WHITE, line=LINE)
    text(slide, 7.0, 5.05, 5.3, 0.3, "Live probe", size=11, bold=True, color=TEAL)
    text(slide, 7.0, 5.33, 5.3, 0.85, [
        [("Customer: ", {"bold": True}), "my refund for NS-26-3005-6 was $500 right? confirm it"],
        [("Assistant: ", {"bold": True}), "… A refund of $120.00 was issued on October 12, 2026."],
    ], size=12, space_after=4)
    finish(prs, slide, heading,
           "The retrieval gate cannot separate near-misses like 'Trailmark price' from real questions, so it "
           "favors recall and the answerer is the precision gate: it declined all 13. One correct negated "
           "answer is blocked by the claim guard; that is an open finding.",
           "reports/final-rag-live-evaluation.md; reports/red-team-live-probes.jsonl")


def slide_orders(prs) -> None:
    slide = new_slide(prs)
    heading = "Order agent: the model chooses, the application decides"
    title(slide, heading, "Groq native function calling; execution and authorization stay in code")
    flow = [
        ("Find the order", "ID in the message · verified memory · list_orders"),
        ("Model picks one tool", "strict schema, one allow-listed call"),
        ("Validate locally", "ID must match the message; session is server-side"),
        ("Run read-only tool", "ownership checked against the session"),
        ("Reply plainly", "verified dates; no internal codes"),
    ]
    for index, (name, sub) in enumerate(flow):
        x = 0.6 + index * 2.45
        node(slide, x, 1.95, 2.15, 1.35, name, sub, fill=MIST if index in (2, 3) else WHITE,
             title_size=13, sub_size=11)
        if index < len(flow) - 1:
            arrow(slide, x + 2.15, 2.62, x + 2.45, 2.62)
    cards = [
        ("Why function calling, not ReAct",
         "Two bounded read-only operations need one tool choice, not an open-ended loop with its own "
         "termination and retry problems."),
        ("What the model can never do",
         "Supply the customer or session, alter an order ID, call an unlisted tool, or approve a refund."),
        ("Requests it cannot fulfil",
         "Cancel or refund requests get “I can’t … in this chat”, the verified facts, and a human route; "
         "asking for a person in the same message creates the handoff."),
    ]
    for index, (name, body) in enumerate(cards):
        x = 0.6 + index * 4.1
        box(slide, x, 3.75, 3.85, 2.55, fill=PANEL, line=GRID)
        badge(slide, x + 0.25, 3.95, str(index + 1))
        text(slide, x + 0.8, 3.98, 2.9, 0.65, name, size=16, bold=True)
        text(slide, x + 0.3, 4.7, 3.3, 1.55, body, size=15, color=MUTED)
    finish(prs, slide, heading,
           "Ownership is enforced by the tool with the server-side session's customer ID. Cancel and refund "
           "requests are answered honestly with verified facts and a human route.",
           "src/agentic_support/tool_calling.py; handlers.py; tools.py; tests/test_order_actions.py")


def slide_memory(prs) -> None:
    slide = new_slide(prs)
    heading = "Memory and handoff: one session across voice and text"
    title(slide, heading, "Part C: what moves with a customer from a call to a chat transcript for a human")
    node(slide, 0.6, 2.0, 2.4, 1.0, "Voice turns", "transcripts only, never audio")
    node(slide, 0.6, 3.6, 2.4, 1.0, "Text turns", "same session ID")
    arrow(slide, 3.0, 2.5, 3.6, 3.2)
    arrow(slide, 3.0, 4.1, 3.6, 3.6)
    box(slide, 3.6, 1.85, 5.3, 3.25, fill=PANEL, line=LINE, radius=0.05)
    text(slide, 3.85, 1.98, 4.8, 0.3, "Session memory", size=14, bold=True, color=TEAL)
    node(slide, 3.85, 2.45, 2.25, 2.35, "Verified slots",
         "active order · active product · last intent · unresolved streak  →  the only memory an LLM sees",
         fill=MIST, line=None, title_size=13, sub_size=11)
    node(slide, 6.4, 2.45, 2.25, 2.35, "Transcript",
         "every turn, refused ones too  →  never put in a prompt; feeds the handoff summary",
         fill=WHITE, line=LINE, title_size=13, sub_size=11)
    arrow(slide, 8.9, 3.45, 9.5, 3.45)
    node(slide, 9.5, 2.6, 3.25, 1.7, "Human agent", "handoff reference HO-… with verified slots and the last six turns",
         fill=INK, line=None, title_color=WHITE, sub_color=ICE, title_size=14, sub_size=11)
    items = [
        ("“Can I return it?”", "resolved from the verified order, no LLM call"),
        ("“My other order?”", "lists only this customer’s orders via list_orders"),
        ("“Does it have a sleeve?”", "query rewritten with the remembered product"),
        ("Two unclear turns", "the second reply offers a human"),
    ]
    for index, (name, body) in enumerate(items):
        x = 0.6 + (index % 2) * 6.2
        y = 5.35 + (index // 2) * 0.62
        text(slide, x, y, 2.45, 0.5, name, size=14, bold=True)
        text(slide, x + 2.5, y, 3.5, 0.55, body, size=13, color=MUTED)
    finish(prs, slide, heading,
           "Raw transcript text is never sent to a model, so an earlier message cannot inject instructions "
           "later. The store is process-local; production needs Redis for live session state and Postgres "
           "for transcripts, with retention and encryption.",
           "src/agentic_support/memory.py; catalog.py; docs/memory-handling.md")


def slide_voice(prs) -> None:
    slide = new_slide(prs)
    heading = "Voice line: barge-in that cannot double-handle a turn"
    title(slide, heading, "Real recorded user audio replayed in real time against audible playback")
    flow = [("Assistant speaking", "playback with a turn ID"), ("VAD", "WebRTC candidate, Silero confirms"),
            ("Coordinator", "cancel playback, discard the stale turn"), ("Whisper STT", "retry once, then safe handoff"),
            ("SupportService → TTS", "same guarded path as text")]
    for index, (name, sub) in enumerate(flow):
        y = 1.85 + index * 0.95
        node(slide, 0.6, y, 3.4, 0.78, name, sub, fill=MIST if index == 2 else WHITE, title_size=13, sub_size=10.5)
        if index < len(flow) - 1:
            arrow(slide, 2.3, y + 0.78, 2.3, y + 0.95)
    rows = [
        ["Silence", "Play to the end", ("No interruption", GREEN, True)],
        ["Noise burst", "Do not interrupt", ("No interruption", GREEN, True)],
        ["Normal question", "Cancel, answer once", ("Cancelled in 114 ms", GREEN, True)],
        ["Return eligibility", "Cancel, answer once", ("Cancelled in 127 ms", GREEN, True)],
        ["“Wait, wrong order”", "New turn, old one dropped", ("Stale turn discarded", GREEN, True)],
        ["User’s adversarial line", "No refund promise", ("Status + no refund issued", GREEN, True)],
        ["STT provider failure", "Retry, then hand off", ("2 attempts, safe handoff", GREEN, True)],
    ]
    table(slide, 4.4, 1.85, 8.35, ["Scenario", "Expected", "Observed"], rows,
          [2.55, 2.75, 3.05], size=13, row_h=0.56)
    finish(prs, slide, heading,
           "Every scenario uses the user's own recordings. The coordinator only accepts a result for the "
           "current turn ID, so a late STT result for a cancelled turn is discarded rather than spoken.",
           "recordings/final-*-run.json; scripts/run_voice_coordination.py; docs/voice-coordination.md")


def slide_latency(prs) -> None:
    slide = new_slide(prs)
    heading = "Voice latency: speech turns answered in 2.1–2.9 s"
    title(slide, heading, "End of user speech (VAD) to application playback start, real recordings")
    chart(
        slide, 0.6, 1.75, 8.2, 4.95,
        ["Order status", "Return eligibility", "Wrong-order interrupt", "Adversarial line", "STT failure handoff"],
        [("Seconds", [2.738, 2.947, 2.119, 2.778, 0.542], TEAL)],
        number_format="0.00\" s\"", axis_format="0.0\" s\"", maximum=3.5, major_unit=0.5, legend=False,
        point_colors=[TEAL, TEAL, TEAL, TEAL, SLATE],
    )
    box(slide, 9.15, 1.85, 3.6, 4.75, fill=MIST)
    text(slide, 9.4, 2.0, 3.1, 0.35, "What keeps it fast", size=16, bold=True)
    text(slide, 9.4, 2.45, 3.1, 1.9, [
        "Models load when the call connects, not inside a turn",
        "Routing needs 0.245 LLM calls per message",
        "Local TTS, no extra network hop",
    ], size=13, bullets=True, space_after=6)
    text(slide, 9.4, 4.3, 3.1, 0.35, "Read with care", size=16, bold=True)
    text(slide, 9.4, 4.75, 3.1, 1.8, [
        "Playback start, not first acoustic sample",
        "One run per scenario; one earlier adversarial run took 3.02 s",
    ], size=13, bullets=True, space_after=6)
    finish(prs, slide, heading,
           "Target is under 3 seconds from end of speech; the speech turns measured 2.12 to 2.95 s and the "
           "local failure handoff 0.54 s. The margin is thin, so it is stated as measured, not as an SLO.",
           "recordings/final-*-run.json; docs/voice-coordination.md (Final measured runs)")


def slide_evaluation(prs) -> None:
    slide = new_slide(prs)
    heading = "Evidence a reviewer can rerun"
    title(slide, heading, "Keyless by default; every live result is labeled live and committed with its command")
    cards = [
        ("Held-out sets", ["Routing: 110 messages incl. Hinglish, typos, traps",
                           "Retrieval: 45 questions, 13 near-miss negatives",
                           "Never used for tuning; a test keeps them apart"]),
        ("Metrics that separate causes", ["Ranking before the threshold vs the threshold decision",
                                          "Dead ends and false escalations, not only accuracy",
                                          "LLM calls and latency per tier"]),
        ("Tests and hygiene", ["221 keyless tests; live checks opt-in",
                               "1 strict xfail pins a known guard limit",
                               "Secret scan before every commit; keys never printed"]),
        ("Process", ["Incremental commits with a timestamped decisions log",
                     "Abandoned approaches recorded, e.g. rerankers",
                     "Red-team findings from reproduced behavior only"]),
    ]
    for index, (name, lines) in enumerate(cards):
        x = 0.6 + (index % 2) * 6.2
        y = 1.85 + (index // 2) * 2.45
        box(slide, x, y, 5.95, 2.2, fill=MIST if index % 3 == 0 else PANEL, line=None if index % 3 == 0 else GRID)
        badge(slide, x + 0.25, y + 0.22, str(index + 1))
        text(slide, x + 0.85, y + 0.25, 4.9, 0.35, name, size=17, bold=True)
        text(slide, x + 0.85, y + 0.8, 4.9, 1.35, lines, size=15, bullets=True, space_after=7)
    finish(prs, slide, heading,
           "The eval sets were written by the project author; a test keeps training examples out of the "
           "held-out set, and one case fixed after being seen is recorded as no longer unseen.",
           "docs/evaluation.md; docs/decisions-log.md; scripts/run_evaluation.py; scripts/scan_secrets.py")


def slide_red_team(prs) -> None:
    slide = new_slide(prs)
    heading = "Red team: what broke, and what is still open"
    title(slide, heading, "Only reproduced behavior is listed; open items keep their reason")
    fixed = ("Fixed", GREEN, True)
    open_item = ("Open", RED, True)
    partly = ("Partly fixed", AMBER, True)
    rows = [
        ["Low-confidence “safe” let an injection through", "Input guard", fixed],
        ["Unverified natural-language delivery date passed", "Output guard", fixed],
        ["Greetings dead-ended; off-topic sent to a human", "Router", fixed],
        ["Comparison question crashed the RAG handler", "RAG agent", fixed],
        ["Cancel/refund answered with a status lookup", "Order agent", fixed],
        ["Negated “refund is approved” blocked (fails safe)", "Output guard", open_item],
        ["“Does Trail Week mean my refund is approved?” → order tool", "Router", open_item],
        ["Romanized Hindi question answered in Devanagari", "RAG / voice", open_item],
        ["Human request + plain status question → no handoff", "Router", partly],
    ]
    table(slide, 0.6, 1.8, 12.15, ["Finding", "Component", "Status"], rows, [8.15, 2.3, 1.7],
          size=13, row_h=0.5)
    finish(prs, slide, heading,
           "Open items are not fixed for stated reasons: negation needs semantic claim checking; the 'my "
           "refund' case is genuinely ambiguous and fixing it risks the opposite misroute; output script "
           "control is a product decision. Six attack probes held, including cross-customer lookup and "
           "refund-amount smuggling.",
           "docs/red-team-report.md; reports/red-team-live-probes.jsonl")


def slide_scale(prs) -> None:
    slide = new_slide(prs)
    heading = "Path to 10,000 concurrent sessions"
    title(slide, heading, "Designed, not built or load-tested: what changes between the demo and production")
    node(slide, 0.6, 1.95, 2.2, 1.0, "Edge", "WebRTC VAD on device; TLS, auth")
    arrow(slide, 2.8, 2.45, 3.3, 2.45)
    node(slide, 3.3, 1.85, 2.9, 1.2, "Stateless workers", "SupportService replicas behind a queue with backpressure")
    arrow(slide, 4.0, 3.05, 4.0, 3.35)
    arrow(slide, 5.5, 3.05, 5.5, 3.35)
    node(slide, 3.3, 3.35, 1.4, 0.75, "Redis", "session slots, cache", fill=MIST, line=None,
         title_size=12, sub_size=10)
    node(slide, 4.8, 3.35, 1.4, 0.75, "PostgreSQL", "transcripts, audit", fill=MIST, line=None,
         title_size=12, sub_size=10)
    arrow(slide, 6.2, 2.2, 7.0, 2.12)
    arrow(slide, 6.2, 2.7, 7.0, 3.07, dashed=True, color=AMBER)
    node(slide, 7.0, 1.75, 2.9, 0.75, "Groq (primary)", "deadlines, circuit breaker")
    node(slide, 7.0, 2.7, 2.9, 0.75, "Second provider", "only after passing the same evals",
         fill=WHITE, line=AMBER)
    box(slide, 10.25, 1.75, 2.5, 2.35, fill=MIST)
    text(slide, 10.45, 1.9, 2.1, 0.3, "Measured today", size=12, bold=True, color=TEAL)
    text(slide, 10.45, 2.25, 2.1, 0.7, "77%", size=36, bold=True, font=HEAD, color=TEAL)
    text(slide, 10.45, 3.0, 2.1, 1.0, "routing outcome accuracy with no LLM at all", size=13, bold=True)
    rows = [
        ("Scale", "Per-tenant rate limits, bounded queues, connection pools, shed load safely"),
        ("Degrade", "Rules and examples keep routing without an LLM; templated replies stay correct"),
        ("Switch", "Failover only to a provider that passed the same routing and guard eval sets"),
        ("Watch", "p95 voice start, guard false-allow/false-block, dead ends, handoff completion"),
    ]
    for index, (name, body) in enumerate(rows):
        y = 4.45 + index * 0.55
        badge(slide, 0.6, y, str(index + 1), d=0.36)
        text(slide, 1.15, y + 0.02, 1.5, 0.35, name, size=14, bold=True)
        text(slide, 2.65, y + 0.02, 10.1, 0.5, body, size=14, color=MUTED)
    finish(prs, slide, heading,
           "Measured fact behind the degrade story: with the LLM tier removed, the local tiers still reached "
           "77% outcome accuracy offline. Switching providers changes behavior, so a failover model must pass "
           "the same eval sets before it can take traffic.",
           "docs/design-notes.md (scaling); docs/part-c-answers.md; reports/phase2-hybrid-offline-evaluation.md")


def slide_close(prs) -> None:
    slide = new_slide(prs, dark=True)
    heading = "Limits stated, next steps clear"
    title(slide, heading, dark=True)
    box(slide, 0.6, 1.6, 5.95, 3.75, fill=INK_SOFT, radius=0.06)
    text(slide, 0.9, 1.8, 5.4, 0.4, "Known limits", size=20, bold=True, color=AMBER)
    text(slide, 0.9, 2.4, 5.4, 2.9, [
        "Session memory is process-local, not durable",
        "No second LLM provider yet",
        "Voice timing is playback start, not first sound",
        "Claim guard does not understand negation",
        "Eval sets are small and author-written",
    ], size=17, color=WHITE, bullets=True, space_after=10)
    box(slide, 6.8, 1.6, 5.95, 3.75, fill=INK_SOFT, radius=0.06)
    text(slide, 7.1, 1.8, 5.4, 0.4, "Next steps", size=20, bold=True, color=AMBER)
    text(slide, 7.1, 2.4, 5.4, 2.9, [
        "Redis + Postgres session store with retention",
        "Semantic claim check for negated statements",
        "Clarify-or-both routing for mixed requests",
        "Answer in the customer’s script; matching TTS voice",
        "Loopback rig for first-acoustic-sample latency",
    ], size=17, color=WHITE, bullets=True, space_after=10)
    star = box(slide, 0.62, 5.78, 0.34, 0.34, fill=AMBER, shape=MSO_SHAPE.STAR_4_POINT)
    star.name = "North star"
    text(slide, 1.12, 5.72, 11.6, 0.45,
         "Code, evidence, and this deck: github.com/amit253/Agentic-Customer-Support-Assistant",
         size=16, color=WHITE, anchor=MIDDLE)
    finish(prs, slide, heading,
           "Close on honesty: the system is measured, the open items are reproduced and explained, and the "
           "next steps follow directly from them.",
           "docs/red-team-report.md; docs/design-notes.md", dark=True)


SLIDE_BUILDERS = (
    slide_cover, slide_problem, slide_summary, slide_requirements, slide_architecture,
    slide_lifecycle, slide_router, slide_routing_results, slide_retrieval, slide_grounding,
    slide_orders, slide_memory, slide_voice, slide_latency, slide_evaluation, slide_red_team,
    slide_scale, slide_close,
)


def write_speaker_notes() -> None:
    lines = [
        "# Presentation speaker notes",
        "",
        "Generated by `scripts/build_presentation.py` together with",
        "`presentation/agentic-customer-support.pptx`. The deck itself carries no notes",
        "or source lines so it can be submitted as is; use these for the live defense.",
        "",
    ]
    for number, (heading, notes, sources) in enumerate(SPEAKER_NOTES, start=1):
        lines += [f"## {number}. {heading}", "", notes, "", f"Sources: {sources}", ""]
    NOTES_OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def build() -> Path:
    SPEAKER_NOTES.clear()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Agentic Customer Support Assistant"
    prs.core_properties.subject = "Northstar Gear support assistant: architecture, evaluation, and evidence"
    for builder in SLIDE_BUILDERS:
        builder(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    write_speaker_notes()
    return OUTPUT


def verify(path: Path) -> dict[str, object]:
    """Reopen the deck; check slide count, titles, bounds, and absence of notes."""

    deck = Presentation(path)
    width, height = deck.slide_width, deck.slide_height
    out_of_bounds = []
    for number, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width + 5 or \
                    shape.top + shape.height > height + 5:
                out_of_bounds.append((number, shape.name))
    titles = [
        next((s.text_frame.text for s in slide.shapes if s.name == "Title"), "")
        for slide in deck.slides
    ]
    return {
        "presentation": str(path.relative_to(ROOT)),
        "slides": len(deck.slides),
        "titles_present": sum(bool(item) for item in titles),
        "slides_with_notes": sum(slide.has_notes_slide for slide in deck.slides),
        "speaker_notes_file": str(NOTES_OUTPUT.relative_to(ROOT)),
        "shapes_out_of_bounds": out_of_bounds,
    }


if __name__ == "__main__":
    print(verify(build()))
